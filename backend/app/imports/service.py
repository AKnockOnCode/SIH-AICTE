"""CSV/XLSX import jobs: upload -> sheet -> column mapping -> dry-run validation -> confirmed commit.

Policy (documented): valid rows can be committed while invalid rows are skipped, but only after
the user explicitly acknowledges the number of skipped rows. Each commit re-validates inside one
database transaction; any unexpected error rolls the whole commit back (status FAILED).
Source identifiers are preserved; internal IDs are always generated. Public signals are only
ever imported as PUBLIC_SIGNALS, never as complaints.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import timedelta
from typing import Any

from ..alerts.service import sweep
from ..audit import chain
from ..core import clock
from ..core.config import get_settings
from ..core.errors import bad_request, conflict, forbidden, not_found, unprocessable
from ..core.security import User
from ..core.util import canonical_json, dumps, loads, next_id, safe_filename, sha256_hex
from ..institutions import matcher
from ..institutions import service as inst_service
from ..institutions.normalize import normalize_name, normalize_permanent_id
from ..scoring import classifier, priority
from ..scoring.rules import category_rules
from .spec import FIELDS, ParseError, is_blank, parse_value, suggest_mapping

KINDS = tuple(FIELDS)
MAX_ROWS = 5000


# ----------------------------------------------------------------------------- parsing
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise unprocessable("Could not decode the CSV file (try saving it as UTF-8).")


def _dedupe(headers: list[Any]) -> list[str]:
    out, seen = [], {}
    for i, h in enumerate(headers):
        name = str(h).strip() if not is_blank(h) else f"Column {i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        out.append(name)
    return out


def read_table(path, file_type: str, sheet: str | None) -> tuple[list[str], list[str], list[dict[str, Any]]]:
    """Return (sheets, columns, rows) — rows are dicts keyed by column name; blank rows dropped."""
    if file_type == "csv":
        text = _decode(path.read_bytes())
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
        if not rows:
            raise unprocessable("The file has no rows.")
        header = _dedupe(rows[0])
        data = [{header[i]: (r[i] if i < len(r) else None) for i in range(len(header))} for r in rows[1:]]
        return ["(csv)"], header, data[:MAX_ROWS + 1]
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        raise unprocessable("Could not open the Excel workbook (use .xlsx; old .xls is not supported).") from None
    sheets = wb.sheetnames
    ws = wb[sheet] if sheet and sheet in sheets else wb[sheets[0]]
    raw = [r for r in ws.iter_rows(values_only=True) if any(not is_blank(c) for c in r)]
    wb.close()
    if not raw:
        raise unprocessable("The selected sheet is empty.")
    header = _dedupe(list(raw[0]))
    data = [{header[i]: (r[i] if i < len(r) else None) for i in range(len(header))} for r in raw[1:]]
    return sheets, header, data[:MAX_ROWS + 1]


def _preview(columns: list[str], rows: list[dict[str, Any]]) -> list[dict[str, str]]:
    def show(v: Any) -> str:
        if is_blank(v):
            return ""
        if hasattr(v, "isoformat"):
            return v.isoformat()[:10]
        return str(v)[:120]

    return [{c: show(r.get(c)) for c in columns} for r in rows[:10]]


def _job(db, job_id: str) -> dict[str, Any]:
    job = db.one("SELECT * FROM import_jobs WHERE id = ?", (job_id,))
    if job is None:
        raise not_found("Import job", job_id)
    return job


def _path(job: dict[str, Any]):
    return get_settings().storage_dir / "imports" / safe_filename(job["stored_name"])


def _check_kind_permission(user: User, kind: str) -> None:
    if kind == "PUBLIC_SIGNALS" and not user.has("signals.import"):
        raise forbidden("Importing public signals requires the Competent Authority (signals.import).")


def upload(db, user: User, kind: str, file_name: str, content: bytes) -> dict[str, Any]:
    if kind not in KINDS:
        raise bad_request(f"kind must be one of {KINDS}")
    _check_kind_permission(user, kind)
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if ext not in ("csv", "xlsx"):
        raise unprocessable("Upload a .csv or .xlsx file.")
    limit = get_settings().max_upload_mb * 1024 * 1024
    if not content or len(content) > limit:
        raise unprocessable(f"File is empty or larger than {get_settings().max_upload_mb} MB.")
    with db.tx():
        job_id = next_id(db, "import", "IMP-", 4)
        stored = f"{job_id}.{ext}"
        folder = get_settings().storage_dir / "imports"
        folder.mkdir(parents=True, exist_ok=True)
        (folder / stored).write_bytes(content)
        try:
            sheets, columns, rows = read_table(folder / stored, ext, None)
        except Exception:
            (folder / stored).unlink(missing_ok=True)
            raise
        if len(rows) > MAX_ROWS:
            raise unprocessable(f"At most {MAX_ROWS} rows per import in this prototype.")
        now = clock.now_iso()
        db.insert("import_jobs", {
            "id": job_id, "kind": kind, "status": "UPLOADED", "file_name": safe_filename(file_name), "file_type": ext,
            "stored_name": stored, "file_sha256": hashlib.sha256(content).hexdigest(), "size_bytes": len(content),
            "sheets_json": dumps(sheets), "sheet_name": sheets[0], "columns_json": dumps(columns), "row_count": len(rows),
            "mapping_json": None, "mapping_hash": None, "validation_json": None, "summary_json": None,
            "policy": "SKIP_INVALID_ROWS", "uploaded_by": user.id, "uploaded_role": user.role, "created_at": now,
            "updated_at": now, "validated_at": None, "committed_at": None,
        })
        chain.record(db, user, "IMPORT_UPLOADED", "import_job", job_id,
                     after={"kind": kind, "file_name": safe_filename(file_name), "sha256": hashlib.sha256(content).hexdigest(),
                            "rows": len(rows), "sheets": sheets})
    return describe(db, job_id, columns=columns, rows=rows)


def describe(db, job_id: str, columns: list[str] | None = None, rows: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    job = _job(db, job_id)
    if columns is None or rows is None:
        _s, columns, rows = read_table(_path(job), job["file_type"], job["sheet_name"])
    mapping = loads(job["mapping_json"], None) or suggest_mapping(job["kind"], columns)
    return {"job": public_job(job), "sheets": loads(job["sheets_json"], []), "sheet": job["sheet_name"], "columns": columns,
            "row_count": len(rows), "preview": _preview(columns, rows), "suggested_mapping": mapping,
            "fields": FIELDS[job["kind"]]}


def select_sheet(db, user: User, job_id: str, sheet: str) -> dict[str, Any]:
    with db.tx():
        job = _job(db, job_id)
        _check_kind_permission(user, job["kind"])
        if job["status"] == "COMMITTED":
            raise conflict("This import is already committed.")
        sheets = loads(job["sheets_json"], [])
        if sheet not in sheets:
            raise bad_request(f"Sheet must be one of {sheets}")
        _s, columns, rows = read_table(_path(job), job["file_type"], sheet)
        db.update("import_jobs", "id", job_id, {"sheet_name": sheet, "columns_json": dumps(columns), "row_count": len(rows),
                                                "mapping_json": None, "mapping_hash": None, "status": "UPLOADED",
                                                "updated_at": clock.now_iso()})
    return describe(db, job_id, columns, rows)


def public_job(job: dict[str, Any]) -> dict[str, Any]:
    out = {k: job[k] for k in ("id", "kind", "status", "file_name", "file_type", "file_sha256", "size_bytes", "sheet_name",
                               "row_count", "mapping_hash", "policy", "uploaded_by", "uploaded_role", "created_at",
                               "updated_at", "validated_at", "committed_at")}
    out["mapping"] = loads(job["mapping_json"], None)
    out["validation"] = loads(job["validation_json"], None)
    out["summary"] = loads(job["summary_json"], None)
    return out


# ----------------------------------------------------------------------------- planning
def _context(db) -> dict[str, Any]:
    return {
        "today": clock.today(),
        "states": [r["name"] for r in db.all("SELECT name FROM states")],
        "categories": {k: v["label"] for k, v in category_rules(db).items()},
        "bureaus": {r["code"]: r["name"] for r in db.all("SELECT code, name FROM bureaus")},
        "officers": {r["id"]: r for r in db.all("SELECT id, display_name, bureau_code FROM officers")},
    }


def _mapping_hash(kind: str, sheet: str | None, mapping: dict[str, str]) -> str:
    return sha256_hex(canonical_json({"kind": kind, "sheet": sheet, "mapping": mapping}))[:16]


def _validate_mapping(kind: str, mapping: dict[str, str], columns: list[str]) -> None:
    keys = {f["key"] for f in FIELDS[kind]}
    errors = [f"unknown target field '{k}'" for k in mapping if k not in keys]
    errors += [f"column '{c}' (for {k}) is not in the file" for k, c in mapping.items() if c and c not in columns]
    used = [c for c in mapping.values() if c]
    errors += [f"column '{c}' is mapped twice" for c in sorted({c for c in used if used.count(c) > 1})]
    required = [f["key"] for f in FIELDS[kind] if f["required"]]
    errors += [f"required field '{k}' is not mapped" for k in required if not mapping.get(k)]
    if kind == "COMPLAINTS" and not (mapping.get("institution_name") or mapping.get("permanent_id")):
        errors.append("map 'institution_name' or 'permanent_id'")
    if errors:
        raise unprocessable("Column mapping is invalid.", errors)


def plan(db, kind: str, mapping: dict[str, str], rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Dry run: parse/validate every row and decide its outcome without writing anything."""
    ctx = _context(db)
    fields = {f["key"]: f for f in FIELDS[kind]}
    seen: dict[str, int] = {}
    out = []
    for n, row in enumerate(rows, start=2):  # spreadsheet row numbers (header = row 1)
        issues: list[dict[str, Any]] = []
        vals: dict[str, Any] = {}
        for key, col in mapping.items():
            if not col:
                continue
            raw = row.get(col)
            try:
                vals[key] = parse_value(fields[key]["type"], raw, ctx)
            except ParseError as e:
                sev = "ERROR" if fields[key]["required"] or key in ("received_date", "disposed_date", "category", "bureau",
                                                                       "status", "outcome", "state", "source_url") else "WARNING"
                issues.append({"row_number": n, "field": key, "severity": sev, "code": "INVALID_VALUE", "message": str(e),
                               "value": None if is_blank(raw) else str(raw)[:200]})
                vals[key] = None
        for key, f in fields.items():
            if f["required"] and vals.get(key) is None and not any(i["field"] == key for i in issues):
                issues.append({"row_number": n, "field": key, "severity": "ERROR", "code": "REQUIRED",
                               "message": f"{f['label']} is required", "value": None})
        decision: dict[str, Any] = {"row_number": n, "outcome": "SKIP", "entity_id": None, "notes": None, "values": vals}
        if kind == "COMPLAINTS":
            _plan_complaint(db, vals, issues, decision, seen, n)
        elif kind == "INSTITUTIONS":
            _plan_institution(db, vals, issues, decision, seen, n)
        else:
            _plan_signal(db, vals, issues, decision, seen, n)
        if any(i["severity"] == "ERROR" for i in issues):
            decision["outcome"] = "SKIP"
        decision["issues"] = issues
        out.append(decision)
    return out


def _plan_complaint(db, v, issues, decision, seen, n) -> None:
    if not (v.get("institution_name") or v.get("permanent_id")):
        issues.append({"row_number": n, "field": "institution_name", "severity": "ERROR", "code": "REQUIRED",
                       "message": "Institution name or Permanent ID is required", "value": None})
    sid = v.get("source_case_id")
    if sid:
        if sid in seen:
            issues.append({"row_number": n, "field": "source_case_id", "severity": "ERROR", "code": "DUPLICATE_IN_FILE",
                           "message": f"Same source ID as row {seen[sid]}", "value": sid})
        seen.setdefault(sid, n)
    status = v.get("status") or "RECEIVED"
    if status == "DISPOSED":
        if not v.get("disposed_date") or not v.get("outcome"):
            issues.append({"row_number": n, "field": "outcome", "severity": "ERROR", "code": "DISPOSAL_INCOMPLETE",
                           "message": "Disposed rows need both a disposal date and an outcome", "value": None})
        elif v.get("received_date") and v["disposed_date"] < v["received_date"]:
            issues.append({"row_number": n, "field": "disposed_date", "severity": "ERROR", "code": "DATE_ORDER",
                           "message": "Disposal date is before receipt date", "value": v["disposed_date"].isoformat()})
    if v.get("assigned_officer") and v.get("bureau"):
        off = db.one("SELECT bureau_code FROM officers WHERE id = ?", (v["assigned_officer"],))
        if off and off["bureau_code"] != v["bureau"]:
            issues.append({"row_number": n, "field": "assigned_officer", "severity": "WARNING", "code": "OFFICER_BUREAU",
                           "message": "Officer belongs to another Bureau/Cell; officer will be left unassigned", "value": v["assigned_officer"]})
            v["assigned_officer"] = None
    if not v.get("description"):
        issues.append({"row_number": n, "field": "description", "severity": "WARNING", "code": "NO_DESCRIPTION",
                       "message": "No description: a placeholder will be stored", "value": None})
    if any(i["severity"] == "ERROR" for i in issues):
        return
    existing = db.one("SELECT id FROM complaints WHERE source_ref = ?", (sid,)) if sid else None
    inst_id, conf, note = None, None, None
    pid = normalize_permanent_id(v.get("permanent_id"))
    row_inst = db.one("SELECT id, name FROM institutions WHERE permanent_id = ?", (pid,)) if pid else None
    if row_inst:
        inst_id, conf, note = row_inst["id"], 1.0, f"Permanent ID matched {row_inst['name']}"
    else:
        q = matcher.Query(v.get("institution_name") or "", raw_address=v.get("institution_district"),
                          raw_pincode=v.get("institution_pincode"), raw_state=v.get("institution_state"), raw_permanent_id=pid)
        res = matcher.match(db, q)
        decision["match"] = {"decision": res["decision"], "candidates": res["candidates"][:3]}
        if res["decision"] == "AUTO_LINK":
            top = res["candidates"][0]
            inst_id, conf, note = top["institution_id"], top["confidence"], f"Auto-linked to {top['name']} ({top['confidence']:.0%})"
        else:
            best = res["candidates"][0] if res["candidates"] else None
            if res["decision"] == "REVIEW" and best:
                note = f"Needs match review (best candidate: {best['name']}, {best['confidence']:.0%})"
            else:
                note = "No institution in the Institution Master matches; queued for officer review (link or create)"
            issues.append({"row_number": n, "field": "institution_name", "severity": "WARNING", "code": "MATCH_REVIEW",
                           "message": note, "value": v.get("institution_name")})
    decision.update({"outcome": "UPDATE" if existing else ("INSERT" if inst_id else "INSERT_REVIEW"),
                     "entity_id": existing["id"] if existing else None, "institution_id": inst_id, "confidence": conf,
                     "notes": note})


def _plan_institution(db, v, issues, decision, seen, n) -> None:
    pid = normalize_permanent_id(v.get("permanent_id"))
    key = pid or f"{normalize_name(v.get('name') or '')}|{v.get('state')}"
    if key in seen:
        issues.append({"row_number": n, "field": "permanent_id" if pid else "name", "severity": "ERROR",
                       "code": "DUPLICATE_IN_FILE", "message": f"Duplicate of row {seen[key]}", "value": pid or v.get("name")})
    seen.setdefault(key, n)
    if any(i["severity"] == "ERROR" for i in issues):
        return
    if pid:
        row = db.one("SELECT id, name FROM institutions WHERE permanent_id = ?", (pid,))
        if row:
            decision.update({"outcome": "UPDATE", "entity_id": row["id"], "notes": f"Permanent ID exists ({row['name']})"})
            return
    res = matcher.match(db, matcher.Query(v["name"], raw_address=v.get("district"), raw_pincode=v.get("pincode"),
                                          raw_state=v.get("state"), raw_permanent_id=pid))
    decision["match"] = {"decision": res["decision"], "candidates": res["candidates"][:3]}
    top = res["candidates"][0] if res["candidates"] else None
    if res["decision"] == "AUTO_LINK" and top and not pid:
        decision.update({"outcome": "MATCHED", "entity_id": top["institution_id"],
                         "notes": f"Already in master as {top['name']} ({top['confidence']:.0%}); alias/fields will be added"})
    elif res["decision"] == "REVIEW" and top:
        decision.update({"outcome": "REVIEW", "notes": f"Possible existing institution {top['name']} ({top['confidence']:.0%}); "
                                                      "queued for officer review, not created"})
        issues.append({"row_number": n, "field": "name", "severity": "WARNING", "code": "MATCH_REVIEW",
                       "message": decision["notes"], "value": v["name"]})
    else:
        decision.update({"outcome": "INSERT", "notes": "New institution"})


def _plan_signal(db, v, issues, decision, seen, n) -> None:
    url = v.get("source_url")
    if url and re.search(r"//(www\.)?example\.(org|com|net)\b", url):
        issues.append({"row_number": n, "field": "source_url", "severity": "ERROR", "code": "EXAMPLE_URL",
                       "message": "Example URL from the template: replace it with the real source URL", "value": url})
        return
    if url:
        if url in seen:
            issues.append({"row_number": n, "field": "source_url", "severity": "ERROR", "code": "DUPLICATE_IN_FILE",
                           "message": f"Same URL as row {seen[url]}", "value": url})
        seen.setdefault(url, n)
        if db.one("SELECT id FROM public_signals WHERE source_url = ?", (url,)):
            issues.append({"row_number": n, "field": "source_url", "severity": "ERROR", "code": "ALREADY_IMPORTED",
                           "message": "This source URL is already in Public Signals", "value": url})
    if any(i["severity"] == "ERROR" for i in issues):
        return
    match = None
    if v.get("institution_text"):
        res = matcher.match(db, matcher.Query(v["institution_text"]))
        top = res["candidates"][0] if res["candidates"] else None
        match = {"decision": res["decision"], "institution_id": top["institution_id"] if top else None,
                 "confidence": top["confidence"] if top else None}
    decision.update({"outcome": "INSERT", "notes": "Public signal (stored separately from complaints; unverified by AgniDrishti)",
                     "signal_match": match})


def validate(db, user: User, job_id: str, mapping: dict[str, str]) -> dict[str, Any]:
    job = _job(db, job_id)
    _check_kind_permission(user, job["kind"])
    if job["status"] == "COMMITTED":
        raise conflict("This import is already committed.")
    _s, columns, rows = read_table(_path(job), job["file_type"], job["sheet_name"])
    mapping = {k: v for k, v in mapping.items() if v}
    _validate_mapping(job["kind"], mapping, columns)
    planned = plan(db, job["kind"], mapping, rows)
    counts = _count(planned)
    mhash = _mapping_hash(job["kind"], job["sheet_name"], mapping)
    with db.tx():
        db.run("DELETE FROM import_row_issues WHERE job_id = ?", (job_id,))
        db.many("INSERT INTO import_row_issues(job_id, row_number, field, severity, code, message, value) VALUES (?,?,?,?,?,?,?)",
                [(job_id, i["row_number"], i["field"], i["severity"], i["code"], i["message"], i["value"])
                 for p in planned for i in p["issues"]])
        validation = {"counts": counts, "mapping_hash": mhash, "validated_at": clock.now_iso()}
        db.update("import_jobs", "id", job_id, {"status": "VALIDATED", "mapping_json": dumps(mapping), "mapping_hash": mhash,
                                                "validation_json": dumps(validation), "validated_at": clock.now_iso(),
                                                "updated_at": clock.now_iso()})
        chain.record(db, user, "IMPORT_VALIDATED", "import_job", job_id, after={"mapping": mapping, "counts": counts,
                                                                                "dry_run": True})
    return {"job": public_job(_job(db, job_id)), "counts": counts, "mapping_hash": mhash,
            "rows": [_row_view(p) for p in planned[:60]],
            "issues": [i for p in planned for i in p["issues"]][:200]}


def _row_view(p: dict[str, Any]) -> dict[str, Any]:
    return {"row_number": p["row_number"], "outcome": p["outcome"], "entity_id": p.get("entity_id"),
            "institution_id": p.get("institution_id"), "confidence": p.get("confidence"), "notes": p.get("notes"),
            "errors": sum(1 for i in p["issues"] if i["severity"] == "ERROR"),
            "warnings": sum(1 for i in p["issues"] if i["severity"] == "WARNING")}


def _count(planned: list[dict[str, Any]]) -> dict[str, int]:
    c = {"total": len(planned), "valid": 0, "invalid": 0, "with_warnings": 0, "insert": 0, "insert_needs_review": 0, "update": 0,
         "matched": 0, "review": 0, "skip": 0}
    for p in planned:
        errs = any(i["severity"] == "ERROR" for i in p["issues"])
        c["invalid" if errs else "valid"] += 1
        if any(i["severity"] == "WARNING" for i in p["issues"]):
            c["with_warnings"] += 1
        key = {"INSERT": "insert", "INSERT_REVIEW": "insert_needs_review", "UPDATE": "update", "MATCHED": "matched",
               "REVIEW": "review", "SKIP": "skip"}[p["outcome"]]
        c[key] += 1
    return c


# ----------------------------------------------------------------------------- commit
def commit(db, user: User, job_id: str, mapping_hash: str, acknowledge_skipped: bool) -> dict[str, Any]:
    job = _job(db, job_id)
    _check_kind_permission(user, job["kind"])
    if job["status"] == "COMMITTED":
        raise conflict("This import is already committed.")
    if job["status"] != "VALIDATED" or not job["mapping_json"]:
        raise conflict("Run validation (dry run) before committing.")
    if mapping_hash != job["mapping_hash"]:
        raise conflict("The mapping changed since validation. Validate again, then commit.")
    mapping = loads(job["mapping_json"], {})
    _s, _cols, rows = read_table(_path(job), job["file_type"], job["sheet_name"])
    results: list[dict[str, Any]] = []
    touched_cases: list[str] = []
    try:
        with db.tx():
            planned = plan(db, job["kind"], mapping, rows)  # re-validated inside the transaction
            counts = _count(planned)
            if counts["invalid"] and not acknowledge_skipped:
                raise conflict(f"{counts['invalid']} invalid row(s) would be skipped. Tick 'acknowledge skipped rows' to proceed "
                               "or fix the file.", {"invalid": counts["invalid"]})
            if counts["valid"] == 0:
                raise unprocessable("No valid rows to import.")
            for p in planned:
                if p["outcome"] == "SKIP":
                    results.append({"row": p["row_number"], "outcome": "SKIPPED", "entity_id": None,
                                    "notes": "; ".join(i["message"] for i in p["issues"] if i["severity"] == "ERROR")})
                    continue
                if job["kind"] == "COMPLAINTS":
                    eid, outcome = _commit_complaint(db, user, job, p)
                    touched_cases.append(eid)
                elif job["kind"] == "INSTITUTIONS":
                    eid, outcome = _commit_institution(db, user, job, p)
                else:
                    eid, outcome = _commit_signal(db, user, job, p)
                results.append({"row": p["row_number"], "outcome": outcome, "entity_id": eid, "notes": p.get("notes")})
            if touched_cases:
                priority.recompute_all(db, case_ids=touched_cases)
                from ..conflicts.detector import detect_for_case

                for cid in touched_cases:
                    detect_for_case(db, cid, trigger_type="IMPORT", actor=user)
            summary = {"counts": {k: sum(1 for r in results if r["outcome"] == k) for k in
                                  ("INSERTED", "INSERTED_NEEDS_REVIEW", "UPDATED", "MATCHED", "REVIEW_QUEUED", "SKIPPED")},
                       "total_rows": len(rows), "results": results[:300], "committed_at": clock.now_iso(),
                       "committed_by": user.id, "policy": "SKIP_INVALID_ROWS (acknowledged)" if counts["invalid"] else "ALL_VALID"}
            db.update("import_jobs", "id", job_id, {"status": "COMMITTED", "summary_json": dumps(summary),
                                                    "committed_at": clock.now_iso(), "updated_at": clock.now_iso()})
            chain.record(db, user, "IMPORT_COMMITTED", "import_job", job_id,
                         after={"kind": job["kind"], "file_sha256": job["file_sha256"], "mapping": mapping,
                                "counts": summary["counts"]})
    except Exception as exc:
        from ..core.errors import ApiError

        if not isinstance(exc, ApiError):
            with db.tx():
                db.update("import_jobs", "id", job_id, {"status": "FAILED", "updated_at": clock.now_iso(),
                                                        "summary_json": dumps({"error": "Unexpected error; nothing was imported "
                                                                               "(transaction rolled back)."})})
                chain.record(db, user, "IMPORT_FAILED", "import_job", job_id, remarks=type(exc).__name__)
        raise
    if job["kind"] == "COMPLAINTS":
        sweep(db)
    return {"job": public_job(_job(db, job_id)), "summary": summary}


def _commit_complaint(db, user: User, job: dict[str, Any], p: dict[str, Any]) -> tuple[str, str]:
    from ..workflow.service import append_action

    v = p["values"]
    rules = category_rules(db)
    now = clock.now_iso()
    if p["outcome"] == "UPDATE":
        case = db.one("SELECT * FROM complaints WHERE id = ?", (p["entity_id"],))
        changes: dict[str, Any] = {}
        if v.get("status") and v["status"] != case["status"]:
            changes["status"] = v["status"]
        if v.get("status") == "DISPOSED":
            changes.update({"disposed_date": v["disposed_date"].isoformat(), "outcome": v["outcome"]})
        if v.get("assigned_officer"):
            changes["assigned_officer_id"] = v["assigned_officer"]
        if changes:
            db.update("complaints", "id", case["id"], {**changes, "updated_at": now})
            append_action(db, case, "STATUS_CHANGED" if "status" in changes else "REMARK", user,
                          remarks=f"Updated from import {job['id']} ({job['file_name']})",
                          from_value=case["status"], to_value=changes.get("status"))
        return case["id"], "UPDATED"
    bureau = v["bureau"]
    case_id = next_id(db, f"case:{bureau}", f"{bureau}-", 4, start=10001)
    rule = rules[v["category"]]
    received = v["received_date"]
    inst = db.one("SELECT * FROM institutions WHERE id = ?", (p.get("institution_id"),)) if p.get("institution_id") else None
    status = v.get("status") or "RECEIVED"
    region = inst["region_code"] if inst else (db.val("SELECT region_code FROM states WHERE name = ?", (v.get("institution_state"),))
                                               if v.get("institution_state") else None)
    desc = v.get("description") or f"Imported record from {job['file_name']} — no description provided in the source."
    text = f"{v.get('subject') or rule['label']}. {desc}"
    ml_cat, ml_p = classifier.predict_many([text])[0]
    row = {
        "id": case_id, "source_system": "IMPORT", "source_ref": v.get("source_case_id"), "external_ref": v.get("external_ref"),
        "channel": "IMPORT", "institution_id": inst["id"] if inst else None,
        "institution_name_raw": v.get("institution_name") or (inst["name"] if inst else "(unknown)"),
        "institution_address_raw": v.get("institution_district"), "institution_pincode_raw": v.get("institution_pincode"),
        "permanent_id_raw": v.get("permanent_id"), "match_status": "AUTO_LINKED" if inst else "PENDING_REVIEW",
        "match_confidence": p.get("confidence"), "category_code": v["category"],
        "subject": (v.get("subject") or rule["label"])[:200], "description": desc, "received_date": received.isoformat(),
        "bureau_code": bureau, "origin_bureau_code": bureau, "assigned_officer_id": v.get("assigned_officer"), "status": status,
        "escalation_level": 1 if status == "ESCALATED" else 0, "transfer_count": 0,
        "complainant_count": v.get("complainant_count") or 1, "complainant_type": "OTHER", "evidence_count": v.get("evidence") or 0,
        "evidence_note": None, "reference_flag": v.get("reference_flag") or "NONE", "reference_note": None,
        "sla_days": int(rule["sla_days"]), "expected_disposal_date": (received + timedelta(days=int(rule["sla_days"]))).isoformat(),
        "disposed_date": v["disposed_date"].isoformat() if status == "DISPOSED" else None,
        "outcome": v.get("outcome") if status == "DISPOSED" else None, "outcome_remarks": None, "last_action_type": None,
        "last_action_at": None, "last_action_summary": None, "next_action": None if status == "DISPOSED" else "Review imported case",
        "next_action_due": None, "priority_score": None, "priority_band": None, "priority_computed_at": None,
        "priority_override_band": None, "priority_override_reason": None, "ml_category": ml_cat, "ml_confidence": ml_p,
        "master_case_id": None, "region_code": region, "state": inst["state"] if inst else v.get("institution_state"),
        "is_synthetic": 0, "demo_tag": None, "created_at": now, "updated_at": now, "created_by": user.id,
    }
    db.insert("complaints", row)
    append_action(db, row, "CREATED", user, remarks=f"Imported from {job['file_name']} (job {job['id']}, row {p['row_number']}); "
                  f"source ID {v.get('source_case_id') or 'none'}.")
    if not inst:
        q = matcher.Query(row["institution_name_raw"], raw_address=v.get("institution_district"),
                          raw_pincode=v.get("institution_pincode"), raw_state=v.get("institution_state"),
                          raw_permanent_id=v.get("permanent_id"))
        inst_service.record_review(db, q, matcher.match(db, q), "IMPORT", case_id, status="PENDING")
        return case_id, "INSERTED_NEEDS_REVIEW"
    return case_id, "INSERTED"


def _commit_institution(db, user: User, job: dict[str, Any], p: dict[str, Any]) -> tuple[str | None, str]:
    v = p["values"]
    now = clock.now_iso()
    if p["outcome"] == "REVIEW":
        q = matcher.Query(v["name"], raw_address=v.get("district"), raw_pincode=v.get("pincode"), raw_state=v.get("state"),
                          raw_permanent_id=v.get("permanent_id"))
        rid = inst_service.record_review(db, q, matcher.match(db, q), "IMPORT", job["id"], status="PENDING")
        return rid, "REVIEW_QUEUED"
    trust_id = None
    if v.get("trust_name"):
        t = db.one("SELECT id FROM trusts WHERE lower(name) = lower(?)", (v["trust_name"],))
        if t:
            trust_id = t["id"]
        else:
            trust_id = next_id(db, "trust", "TRS-I", 3)
            db.insert("trusts", {"id": trust_id, "name": v["trust_name"], "registration_ref": None, "state": v.get("state"),
                                 "provenance": f"IMPORTED from {job['file_name']} (job {job['id']})", "is_synthetic": 0,
                                 "notes": v.get("source_note"), "created_at": now})
    if p["outcome"] in ("UPDATE", "MATCHED"):
        inst = db.one("SELECT * FROM institutions WHERE id = ?", (p["entity_id"],))
        changes = {}
        for key, col in (("district", "district"), ("address", "address_line"), ("pincode", "pincode"),
                         ("institution_type", "institution_type"), ("established_year", "established_year")):
            if v.get(key) and not inst[col]:
                changes[col] = v[key]
        if v.get("intake_total") is not None:
            changes.update({"intake_total": v["intake_total"], "intake_source": "IMPORTED"})
        if v.get("courses"):
            changes["courses_json"] = dumps(v["courses"])
        if trust_id and not inst["trust_id"]:
            changes.update({"trust_id": trust_id, "trust_relation_source": "IMPORTED"})
        if changes:
            db.update("institutions", "id", inst["id"], {**changes, "updated_at": now})
        for alias in [v["name"], *(v.get("aliases") or [])]:
            an = normalize_name(alias)
            if an and an != inst["name_norm"]:
                db.run("INSERT OR IGNORE INTO institution_aliases(institution_id, alias, alias_norm, alias_type, source, created_by, "
                       "created_at) VALUES (?,?,?,?,?,?,?)", (inst["id"], alias, an, "IMPORTED", job["id"], user.id, now))
        chain.record(db, user, "INSTITUTION_UPDATED", "institution", inst["id"], after={"changes": list(changes), "import": job["id"]})
        return inst["id"], "UPDATED" if p["outcome"] == "UPDATE" else "MATCHED"
    row = inst_service.create_institution(db, user, {
        "name": v["name"], "state": v["state"], "district": v.get("district"), "pincode": v.get("pincode"),
        "permanent_id": v.get("permanent_id"), "address": v.get("address"), "institution_type": v.get("institution_type"),
        "established_year": v.get("established_year"), "intake_total": v.get("intake_total"),
        "intake_source": "IMPORTED" if v.get("intake_total") is not None else None, "courses": v.get("courses"),
        "trust_id": trust_id, "trust_relation_source": "IMPORTED" if trust_id else None, "is_synthetic": False,
        "provenance": {"source": "IMPORT", "job": job["id"], "file": job["file_name"], "file_sha256": job["file_sha256"],
                       "row": p["row_number"], "note": v.get("source_note")},
    }, source=f"IMPORT:{job['id']}")
    for alias in v.get("aliases") or []:
        an = normalize_name(alias)
        if an and an != row["name_norm"]:
            db.run("INSERT OR IGNORE INTO institution_aliases(institution_id, alias, alias_norm, alias_type, source, created_by, "
                   "created_at) VALUES (?,?,?,?,?,?,?)", (row["id"], alias, an, "IMPORTED", job["id"], user.id, now))
    return row["id"], "INSERTED"


def _commit_signal(db, user: User, job: dict[str, Any], p: dict[str, Any]) -> tuple[str, str]:
    v = p["values"]
    sig_id = next_id(db, "signal", "SIG-", 4)
    m = p.get("signal_match") or {}
    # Same routing as complaint intake: only an unambiguous high-confidence match is linked; a
    # plausible one is kept as UNCERTAIN (shown with its confidence), never silently attached.
    status = {"AUTO_LINK": "MATCHED", "REVIEW": "UNCERTAIN"}.get(m.get("decision") or "", "NO_MATCH")
    db.insert("public_signals", {
        "id": sig_id, "signal_type": v.get("signal_type") or "NEWS_REPORT", "verification": "UNVERIFIED",
        "scope": "INSTITUTION" if v.get("institution_text") else "SECTOR", "title": v["title"], "publisher": v["publisher"],
        "source_url": v["source_url"], "published_date": v["published_date"].isoformat() if v.get("published_date") else None,
        "retrieved_at": None, "summary": v["summary"][:1000], "raw_institution_text": v.get("institution_text"),
        "institution_id": m.get("institution_id") if status != "NO_MATCH" else None, "match_confidence": m.get("confidence"),
        "match_status": status, "notes": f"Imported from {job['file_name']} (job {job['id']}); not verified by AgniDrishti.",
        "import_job_id": job["id"], "created_at": clock.now_iso(),
    })
    return sig_id, "INSERTED"


# ----------------------------------------------------------------------------- templates
def template_bytes(kind: str, fmt: str) -> tuple[bytes, str]:
    fields = FIELDS[kind]
    examples = {
        "INSTITUTIONS": {"name": "Example Institute of Technology", "permanent_id": "1-1234567890", "state": "Maharashtra",
                         "district": "Pune", "pincode": "411001", "intake_total": "540", "courses": "B.Tech (CSE); MBA",
                         "trust_name": "Example Education Society", "source_note": "AICTE dashboard export, 2026"},
        "COMPLAINTS": {"source_case_id": "PGRC/2026/01234", "institution_name": "Example Institute of Technology",
                       "permanent_id": "1-1234567890", "institution_state": "Maharashtra", "category": "Fee fraud / excess fee",
                       "received_date": "27-10-2026", "bureau": "PGRC", "status": "Under process",
                       "subject": "Excess fee collected", "description": "Students report collection of fees above the approved "
                       "structure without receipts.", "complainant_count": "3", "evidence": "yes", "reference_flag": "NONE"},
        "PUBLIC_SIGNALS": {"title": "Public notice on unapproved institutions", "publisher": "AICTE",
                           "source_url": "https://example.org/notice.pdf", "published_date": "01-11-2026",
                           "summary": "Short factual summary in your own words.", "signal_type": "OFFICIAL_NOTICE"},
    }[kind]
    labels = [f["label"] for f in fields]
    example = [examples.get(f["key"], "") for f in fields]
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(labels)
        w.writerow(example)
        return buf.getvalue().encode("utf-8-sig"), "text/csv"
    import xlsxwriter

    out = io.BytesIO()
    wb = xlsxwriter.Workbook(out, {"in_memory": True})
    head = wb.add_format({"bold": True, "bg_color": "#1F3864", "font_color": "#FFFFFF"})
    ws = wb.add_worksheet("Data")
    for j, (lab, ex) in enumerate(zip(labels, example, strict=True)):
        ws.write(0, j, lab, head)
        ws.write(1, j, ex)
        ws.set_column(j, j, max(14, len(lab) + 2))
    dd = wb.add_worksheet("Data dictionary")
    for j, h in enumerate(["Field", "Required", "Type", "Accepted header names", "Notes"]):
        dd.write(0, j, h, head)
    for i, f in enumerate(fields, start=1):
        dd.write_row(i, 0, [f["label"], "yes" if f["required"] else "no", f["type"], ", ".join(f["synonyms"]),
                            f.get("description", "")])
    dd.set_column(0, 0, 32)
    dd.set_column(3, 4, 60)
    wb.close()
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
